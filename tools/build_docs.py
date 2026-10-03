"""Convert the WSQ LP/LG of TGS-2025053207 into the one-day non-WSQ C171 versions."""
import copy, re
from docx import Document

SRC = "source-wsq/courseware/"
OUT = "courseware/"
TITLE = "Git and Github Masterclass"
CODE = "C171"
DROPPED = [3, 9, 10, 11, 147, 148]          # 1-based source slides removed from the deck


def remap(n):
    return n - sum(1 for d in DROPPED if d < n)


def put(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text); return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def drop(p):
    p._p.getparent().remove(p._p)


def cover_and_version(d, summary):
    for p in d.paragraphs[:15]:
        t = p.text.strip()
        if t == "GitHub Foundations Certification Training": put(p, TITLE)
        elif t.startswith("TGS Ref No"): put(p, f"Course Code: {CODE}")
        elif t == "Version 2.2": put(p, "Version 3.0")
    vt = d.tables[0]
    new = copy.deepcopy(vt.rows[-1]._tr)
    vt._tbl.append(new)
    row = vt.rows[-1]
    for c, v in zip(row.cells, ["3.0", "4 October 2026", summary, "Dr Alfred Ang"]):
        put(c.paragraphs[0], v)


# ------------------------------------------------------------------ LP ----
d = Document(SRC + "LP-Github-Foundations-Certification-Training.docx")
cover_and_version(d, "One-day edition for Git and Github Masterclass (C171): all seven domains and fourteen labs "
                     "retimed to 9:30 am – 5:30 pm.")
P = d.paragraphs
for p in list(P):
    t = p.text.strip()
    if t.startswith("Assessment Design") or t.startswith("Day 2 —"):
        drop(p)
    elif t.startswith("Day 1 —"):
        if "\t" in p.text:      # TOC entry
            put(p, "Day 1 · 9:30am – 5:30pm — Git Foundations, Collaboration, Automation and Governance\t4")
        else:
            put(p, "Day 1 · 9:30am – 5:30pm — Git Foundations, Collaboration, Automation and Governance")
    elif t.startswith("Total instructional time"):
        put(p, "Total instructional time: 450 minutes (7.5 instructional hours, 9:30am – 5:30pm), with a single 30-minute lunch break.")
# Assessment Design bullets + second Day total (now orphaned)
P = d.paragraphs
for p in list(P):
    t = p.text.strip()
    if t.startswith(("Written Assessment", "Case Study (CS)", "The instrument type", "Briefing precedes")):
        drop(p)
tots = [p for p in d.paragraphs if p.text.startswith("Total instructional time")]
for p in tots[1:]:
    drop(p)

info = d.tables[1]
for r in list(info.rows):
    k = r.cells[0].text.strip()
    if k == "Course Title": put(r.cells[1].paragraphs[0], TITLE)
    elif k == "Course Code": put(r.cells[1].paragraphs[0], CODE)
    elif k == "Duration": put(r.cells[1].paragraphs[0], "1 day · 7.5 instructional hours")
    elif k == "Daily Timing": put(r.cells[1].paragraphs[0], "9:30 am – 5:30 pm (30-minute lunch break)")
    elif k == "Assessment": info._tbl.remove(r._tr)

# One-day schedule: every WSQ teaching row kept, rescaled proportionally onto 450 min.
items = [
    ("Admin", "Welcome, trainer profile, ground rules and course portal", 30, "2–8"),
    ("Concept", "GH-900 blueprint, Git areas, file states and toolchain", 45, "—"),
    ("Lab", "Lab 01 — Install Git, Git Bash and GitHub Desktop", 45, 31),
    ("Lab", "Lab 02 — Create and Clone a Training Repository", 45, 33),
    ("Lab", "Lab 03 — Status, Diff, Add, Commit and Log", 60, 59),
    ("Concept", "Domain 2 — branch pointers, merge strategies and remotes", 30, 67),
    ("Lab", "Lab 04 — Create Branches and Merge Changes", 60, 61),
    ("Lab", "Lab 05 — Synchronise with Fetch, Pull and Push", 45, 63),
    ("Lab", "Lab 06 — Resolve a Merge Conflict and Recover Safely", 45, 65),
    ("Concept", "Domain 3 — pull requests, issues and collaborative change", 30, 82),
    ("Lab", "Lab 07 — Open a Pull Request and Complete Code Review", 45, 78),
    ("Lab", "Lab 08 — Coordinate Work with Issues, Discussions and Pages", 45, 80),
    ("Concept", "Domain 4 — Actions, Copilot, Codespaces and github.dev", 30, 98),
    ("Lab", "Lab 09 — Build and Diagnose a GitHub Actions Check", 45, 94),
    ("Lab", "Lab 10 — Compare Codespaces, github.dev and Copilot", 45, 96),
    ("Concept", "Domain 5 — GitHub Projects and release coordination", 30, 109),
    ("Lab", "Lab 11 — Plan and Track a Release with GitHub Projects", 45, 107),
    ("Concept", "Domain 6 — privacy, security and administration", 30, 122),
    ("Lab", "Lab 12 — Secure and Administer a Repository", 45, 120),
    ("Lab", "Lab 13 — Contribute through a Fork and Explore the Community", 30, 131),
    ("Lab", "Lab 14 — Build a GH-900 Readiness Plan", 30, 133),
    ("Recap", "Course recap, retrieval practice, Q&A and course feedback", 15, "137–144"),
]
tot = sum(i[2] for i in items)
exact = [i[2] / tot * 90 for i in items]          # 90 five-minute units = 450 min
units = [int(e) for e in exact]
for k in sorted(range(len(items)), key=lambda k: exact[k] - units[k], reverse=True)[:90 - sum(units)]:
    units[k] += 1                                      # largest remainder keeps proportions
mins = [u * 5 for u in units]
assert sum(mins) == 450

sched = d.tables[2]
tmpl = copy.deepcopy(sched.rows[2]._tr)      # a body row
for r in list(sched.rows)[1:]:
    sched._tbl.remove(r._tr)


def hm(m):
    return f"{m // 60}:{m % 60:02d}"


def add(start, n, typ, text, ref):
    tr = copy.deepcopy(tmpl)
    sched._tbl.append(tr)
    row = sched.rows[-1]
    for c, v in zip(row.cells, [f"{hm(start)}–{hm(start + n)}", str(n), typ, text, str(ref)]):
        put(c.paragraphs[0], v)


t = 9 * 60 + 30
lunch_done = False
for (typ, text, _, ref), n in zip(items, mins):
    if not lunch_done and t >= 12 * 60 + 30:
        add(t, 30, "Lunch", "Lunch break", "—"); t += 30; lunch_done = True
    ref = f"{remap(ref)}" if isinstance(ref, int) else ref
    add(t, n, typ, text, ref); t += n
assert t == 17 * 60 + 30, t
sched_rows = [[c.text for c in r.cells] for r in sched.rows]

# second (Day 2) schedule table goes
d.tables[3]._tbl.getparent().remove(d.tables[3]._tbl)
align = d.tables[3]
for r in list(align.rows)[1:]:
    put(r.cells[3].paragraphs[0], str(remap(int(r.cells[3].text))))
d.save(OUT + f"LP-{TITLE} ({CODE}).docx")
for r in sched_rows: print(" | ".join(r))

# ------------------------------------------------------------------ LG ----
d = Document(SRC + "LG-Github-Foundations-Certification-Training.docx")
cover_and_version(d, "One-day edition for Git and Github Masterclass (C171).")
for p in d.paragraphs:
    t = p.text
    if t.startswith("This 16-hour, two-day WSQ course"):
        put(p, "This one-day, 7.5-hour course prepares learners to use Git and GitHub confidently and to revise for the "
               "GH-900 GitHub Foundations certification. The slides are concept-led; this Learner Guide contains the "
               "detailed procedures and evidence requirements.")
    elif "screenshot, assessment or Excel workbook" in t:
        for r in p.runs:
            r.text = r.text.replace("screenshot, assessment or Excel workbook", "screenshot or Excel workbook")
d.save(OUT + f"LG-{TITLE} ({CODE}).docx")
print("ok")
